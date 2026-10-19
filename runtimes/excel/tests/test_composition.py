"""Oracle littéral indépendant ; classeurs fictifs construits pendant le test."""

from hashlib import sha256
from io import BytesIO

import openpyxl
import polars as pl
import pytest

from memlia_excel.tables import apply_operation
from memlia_excel.xlsx import read_table, write_table


@pytest.mark.parametrize(
    "identifier,amount,group",
    [("item", "cost", "category"), ("ref", "quantity", "area")],
)
def test_generic_pipeline_reuses_same_operators_with_new_column_names(
    identifier, amount, group
):
    book = openpyxl.Workbook()
    book.active.title = "Source"
    for row in [
        [identifier, amount],
        ["b", "0.10"],
        ["a", None],
        ["b", "0.20"],
        ["c", "0"],
    ]:
        book.active.append(row)
    stream = BytesIO()
    book.save(stream)
    book.close()
    source = stream.getvalue()
    digest = sha256(source).hexdigest()
    frame = read_table(
        source,
        expected_digest=digest,
        sheet="Source",
        columns={identifier: "string", amount: "decimal"},
        precision=18,
        scale=2,
    )
    lookup = pl.DataFrame(
        {identifier: ["a", "b", "c"], group: ["missing", "present", "zero"]}
    )
    operations = [
        {
            "op": "join",
            "leftKeys": [identifier],
            "rightKeys": [identifier],
            "how": "left",
            "cardinality": "many_to_one",
            "onUnmatched": "reject",
        },
        {
            "op": "aggregate",
            "groupBy": [group],
            "values": [
                {"column": amount, "as": "total", "function": "sum", "allNull": "null"}
            ],
        },
        {
            "op": "sort",
            "keys": [{"column": group, "direction": "asc", "nulls": "last"}],
        },
    ]

    def calculate():
        result = frame
        for op in operations:
            result = apply_operation(result, op, precision=18, scale=2, right=lookup)
        return result

    first, replay = calculate(), calculate()
    assert first.equals(replay)
    artifact = write_table(first, sheet="Result")
    result_book = openpyxl.load_workbook(BytesIO(artifact), data_only=False)
    assert list(result_book.active.values) == [
        (group, "total"),
        ("missing", None),
        ("present", "0.30"),
        ("zero", "0.00"),
    ]
    result_book.close()
    assert sha256(source).hexdigest() == digest
