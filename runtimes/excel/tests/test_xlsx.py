from hashlib import sha256
from io import BytesIO
from decimal import Decimal
from zipfile import ZipFile, ZIP_DEFLATED

import openpyxl
import polars as pl
import pytest

from memlia_excel import xlsx
from memlia_excel.tables import KernelError


def workbook(rows):
    book = openpyxl.Workbook()
    book.active.title = "Data"
    for row in rows:
        book.active.append(row)
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def read(data, **kwargs):
    return xlsx.read_table(
        data,
        expected_digest=sha256(data).hexdigest(),
        sheet="Data",
        columns={"id": "string", "amount": "decimal"},
        precision=18,
        scale=2,
        **kwargs,
    )


def test_real_xlsx_preserves_missing_zero_and_exact_decimal_without_mutating_source(
    tmp_path,
):
    source = tmp_path / "source.xlsx"
    source.write_bytes(
        workbook(
            [["id", "amount"], ["a", None], ["b", "0"], ["c", "9999999999999999.99"]]
        )
    )
    before = source.read_bytes()
    table = read(before)
    assert table.rows() == [
        ("a", None),
        ("b", Decimal("0")),
        ("c", Decimal("9999999999999999.99")),
    ]
    output = xlsx.write_table(table, sheet="Result")
    inspected = openpyxl.load_workbook(BytesIO(output), data_only=False)
    assert list(inspected.active.values) == [
        ("id", "amount"),
        ("a", None),
        ("b", "0.00"),
        ("c", "9999999999999999.99"),
    ]
    inspected.close()
    assert source.read_bytes() == before


def test_changed_input_digest_is_rejected():
    data = workbook([["id", "amount"], ["a", "1"]])
    with pytest.raises(KernelError, match="source_changed"):
        xlsx.read_table(
            data,
            expected_digest="0" * 64,
            sheet="Data",
            columns={"id": "string"},
            precision=18,
            scale=2,
        )


@pytest.mark.parametrize(
    "formula",
    ["=1+1", '=HYPERLINK("https://example.invalid","x")', "='[other.xlsx]Sheet1'!A1"],
)
def test_actual_formula_is_rejected_without_using_cached_value(formula):
    with pytest.raises(KernelError, match="formula"):
        read(workbook([["id", "amount"], [formula, "1"]]))


@pytest.mark.parametrize(
    "text", ["=1+1", "+SUM(A1)", "-2+3", "@SUM(A1)", "=cmd|' /C harmless'!A0"]
)
def test_formula_looking_strings_stay_literal_after_real_save_and_readback(text):
    result = xlsx.write_table(pl.DataFrame({"value": [text]}), sheet="Result")
    book = openpyxl.load_workbook(BytesIO(result), data_only=False)
    assert book.active["A2"].value == text
    assert book.active["A2"].data_type == "s"
    book.close()


def test_long_text_is_rejected_before_openpyxl_can_truncate():
    with pytest.raises(KernelError, match="text_length"):
        xlsx.write_table(pl.DataFrame({"value": ["x" * 32768]}), sheet="Result")


def test_numeric_float_is_rejected_as_inexact_decimal():
    with pytest.raises(KernelError, match="decimal"):
        read(workbook([["id", "amount"], ["a", 0.1]]))


def test_cell_limit_before_reading_large_sheet():
    with pytest.raises(KernelError, match="cell_limit"):
        read(workbook([["id", "amount"], ["a", "1"]]), max_cells=3)


def altered_archive(name, content):
    original = workbook([["id", "amount"], ["a", "1"]])
    output = BytesIO()
    with (
        ZipFile(BytesIO(original)) as archive,
        ZipFile(output, "w", ZIP_DEFLATED) as target,
    ):
        for item in archive.infolist():
            if item.filename != name:
                target.writestr(item.filename, archive.read(item))
        target.writestr(name, content)
    return output.getvalue()


@pytest.mark.parametrize(
    "name,content",
    [
        ("../outside", b"x"),
        ("xl/vbaProject.bin", b"x"),
        ("xl/externalLinks/externalLink1.xml", b"<x/>"),
        (
            "xl/worksheets/_rels/sheet1.xml.rels",
            b'<Relationships><Relationship TargetMode="External" Target="https://example.invalid"/></Relationships>',
        ),
        ("docProps/core.xml", b'<!DOCTYPE x [<!ENTITY a "x">]><x>&a;</x>'),
    ],
)
def test_unsafe_or_unqualified_package_is_rejected(name, content):
    with pytest.raises(KernelError):
        read(altered_archive(name, content))


def test_expanded_archive_limit():
    data = altered_archive("docProps/core.xml", b"x" * 20000)
    with pytest.raises(KernelError, match="archive_limit"):
        read(data, max_expanded_bytes=10000)


def test_false_declared_dimension_cannot_hide_real_rows():
    data = workbook([["id", "amount"], ["a", "1"], ["b", "2"]])
    with ZipFile(BytesIO(data)) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").replace(
            b'ref="A1:B3"', b'ref="A1:B1"'
        )
    changed = altered_archive("xl/worksheets/sheet1.xml", xml)
    assert read(changed).rows() == [("a", Decimal("1")), ("b", Decimal("2"))]


def test_readback_mismatch_never_returns_artifact(monkeypatch):
    real_load = openpyxl.load_workbook

    def corrupt_readback(*args, **kwargs):
        book = real_load(*args, **{**kwargs, "read_only": False})
        book.active["A2"] = "incorrect"
        return book

    monkeypatch.setattr(openpyxl, "load_workbook", corrupt_readback)
    with pytest.raises(KernelError, match="readback_mismatch"):
        xlsx.write_table(pl.DataFrame({"value": ["correct"]}), sheet="Result")


@pytest.mark.parametrize(
    "body,expected",
    [
        (
            [["a", "1"], [], ["b", "2"]],
            [("a", Decimal("1.00")), (None, None), ("b", Decimal("2.00"))],
        ),
        ([[], ["b", "2"]], [(None, None), ("b", Decimal("2.00"))]),
    ],
)
def test_missing_rows_are_materialized_as_nulls(body, expected):
    assert read(workbook([["id", "amount"], *body])).rows() == expected


@pytest.mark.parametrize(
    "formula",
    ["1+1", "'[https://example.invalid/other.xlsx]Sheet1'!$A$1", "'Data'!$A$2"],
)
def test_defined_names_are_outside_plain_workbook_profile(formula):
    from openpyxl.workbook.defined_name import DefinedName

    book = openpyxl.load_workbook(BytesIO(workbook([["id", "amount"], ["a", "1"]])))
    book.defined_names.add(DefinedName("synthetic_name", attr_text=formula))
    output = BytesIO()
    book.save(output)
    book.close()
    with pytest.raises(KernelError, match="unsupported_workbook_object"):
        read(output.getvalue())


def test_other_unqualified_workbook_objects_are_rejected():
    data = workbook([["id", "amount"], ["a", "1"]])
    with ZipFile(BytesIO(data)) as archive:
        xml = archive.read("xl/workbook.xml").replace(
            b"</workbook>", b"<customWorkbookViews/><externalReferences/></workbook>"
        )
    with pytest.raises(KernelError, match="unsupported_workbook_object"):
        read(altered_archive("xl/workbook.xml", xml))


@pytest.mark.parametrize(
    "change",
    [
        "duplicate_row",
        "duplicate_cell",
        "reversed_rows",
        "reversed_cells",
        "mismatched_row",
        "missing_row_index",
    ],
)
def test_ambiguous_or_unordered_sheet_coordinates_are_rejected(change):
    from xml.etree import ElementTree as ET

    source = workbook([["id", "amount"], ["a", "1"], ["b", "2"]])
    with ZipFile(BytesIO(source)) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    data = root.find(
        "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheetData"
    )
    if change == "duplicate_row":
        data[2].set("r", "2")
        data[2][0].set("r", "A2")
        data[2][1].set("r", "B2")
    elif change == "duplicate_cell":
        data[1][1].set("r", "A2")
    elif change == "reversed_rows":
        data[:] = [data[0], data[2], data[1]]
    elif change == "reversed_cells":
        data[1][:] = list(reversed(data[1]))
    elif change == "mismatched_row":
        data[1][1].set("r", "B3")
    else:
        del data[1].attrib["r"]
    with pytest.raises(KernelError, match="sheet_coordinates"):
        read(altered_archive("xl/worksheets/sheet1.xml", ET.tostring(root)))


@pytest.mark.parametrize("header", [False, True])
def test_invalid_output_characters_refuse_without_disclosing_cell_content(header):
    import traceback

    private_value = "synthetic\x01example"
    frame = (
        pl.DataFrame({private_value: ["value"]})
        if header
        else pl.DataFrame({"value": [private_value]})
    )
    with pytest.raises(KernelError, match="invalid_output_value") as error:
        xlsx.write_table(frame, sheet="Result")
    assert "synthetic" not in "".join(traceback.format_exception(error.value))
    assert error.value.__cause__ is None
    assert error.value.__context__ is None
