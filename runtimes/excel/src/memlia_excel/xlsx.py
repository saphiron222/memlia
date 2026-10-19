# Copyright (c) 2026 Memlia. All rights reserved.
"""IO XLSX bornée sur octets : aucune résolution de chemin, aucun effet disque."""

from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import re
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import Element

from defusedxml import ElementTree
from defusedxml.common import DefusedXmlException
import openpyxl
from openpyxl.utils.exceptions import IllegalCharacterError
import polars as pl

from .tables import KernelError, decimal_series

_PART = re.compile(
    r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:app|core)\.xml|xl/(?:workbook\.xml|styles\.xml|sharedStrings\.xml|_rels/workbook\.xml\.rels|theme/theme[0-9]+\.xml|worksheets/sheet[0-9]+\.xml))\Z"
)
_SHEET_TAGS = frozenset(
    "worksheet sheetPr outlinePr pageSetUpPr dimension sheetViews sheetView selection sheetFormatPr cols col sheetData row c v is t pageMargins".split()
)

_WORKBOOK_TAGS = frozenset(
    "workbook workbookPr workbookProtection bookViews workbookView sheets sheet definedNames calcPr".split()
)
_MAIN_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _check_sheet_coordinates(root: Element, max_cells: int) -> None:
    """Refuse les coordonnées ambiguës avant le lecteur streaming openpyxl."""
    if root.tag != _MAIN_NS + "worksheet":
        raise KernelError("sheet_coordinates")
    containers = root.findall(_MAIN_NS + "sheetData")
    if len(containers) != 1:
        raise KernelError("sheet_coordinates")
    parents = {child: parent for parent in root.iter() for child in parent}
    for node in root.iter():
        if node.tag == _MAIN_NS + "row" and parents[node] is not containers[0]:
            raise KernelError("sheet_coordinates")
        if node.tag == _MAIN_NS + "c" and parents[node].tag != _MAIN_NS + "row":
            raise KernelError("sheet_coordinates")
    previous_row, max_column = 0, 1
    for row in containers[0]:
        index = row.get("r", "")
        if row.tag != _MAIN_NS + "row" or not re.fullmatch(r"[1-9][0-9]*", index):
            raise KernelError("sheet_coordinates")
        row_number = int(index)
        if row_number <= previous_row or row_number > 1048576:
            raise KernelError("sheet_coordinates")
        previous_row, previous_column = row_number, 0
        for cell in row:
            coordinate = cell.get("r", "")
            if cell.tag != _MAIN_NS + "c" or not re.fullmatch(
                r"[A-Z]{1,3}[1-9][0-9]*", coordinate
            ):
                raise KernelError("sheet_coordinates")
            cell_row, column = openpyxl.utils.cell.coordinate_to_tuple(coordinate)
            if cell_row != row_number or column <= previous_column or column > 16384:
                raise KernelError("sheet_coordinates")
            previous_column = column
            max_column = max(max_column, column)
        if row_number * max_column > max_cells:
            raise KernelError("cell_limit")


def _check_archive(data: bytes, max_expanded_bytes: int, max_cells: int) -> None:
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 512
                or sum(item.file_size for item in entries) > max_expanded_bytes
            ):
                raise KernelError("archive_limit")
            if len({item.filename for item in entries}) != len(entries):
                raise KernelError("archive_duplicate")
            for entry in entries:
                if not _PART.fullmatch(entry.filename) or entry.flag_bits & 1:
                    raise KernelError("unsupported_package_part")
                root = ElementTree.fromstring(archive.read(entry), forbid_dtd=True)
                if entry.filename.startswith("xl/worksheets/"):
                    _check_sheet_coordinates(root, max_cells)
                for node in root.iter():
                    tag = node.tag.rsplit("}", 1)[-1]
                    if entry.filename == "xl/workbook.xml" and (
                        node.tag != _MAIN_NS + tag
                        or tag not in _WORKBOOK_TAGS
                        or (tag == "workbookProtection" and node.attrib)
                    ):
                        raise KernelError("unsupported_workbook_object")
                    if tag == "f":
                        raise KernelError("formula_not_supported")
                    if node.attrib.get("TargetMode", "").lower() == "external":
                        raise KernelError("external_link")
                    if (
                        entry.filename.startswith("xl/worksheets/")
                        and tag not in _SHEET_TAGS
                    ):
                        raise KernelError("unsupported_sheet_object")
    except (BadZipFile, ElementTree.ParseError, DefusedXmlException) as exc:
        raise KernelError("invalid_archive") from exc


def read_table(
    data: bytes,
    *,
    expected_digest: str,
    sheet: str,
    columns: dict[str, str],
    precision: int,
    scale: int,
    max_cells: int = 100000,
    max_expanded_bytes: int = 10000000,
) -> pl.DataFrame:
    """Lit une table rectangulaire avec en-tête exact ; les floats métier sont refusés."""
    if type(data) is not bytes or sha256(data).hexdigest() != expected_digest:
        raise KernelError("source_changed")
    if len(data) > 10000000:
        raise KernelError("archive_limit")
    _check_archive(data, max_expanded_bytes, max_cells)
    book = openpyxl.load_workbook(
        BytesIO(data), read_only=True, data_only=False, keep_links=False
    )
    try:
        if sheet not in book.sheetnames:
            raise KernelError("sheet_missing")
        page = book[sheet]
        page.reset_dimensions()
        rows = []
        cells_seen = 0
        for row in page.values:
            cells_seen += max(1, len(columns))
            if cells_seen > max_cells:
                raise KernelError("cell_limit")
            if len(row) > len(columns):
                raise KernelError("row_width")
            rows.append(tuple(row) + (None,) * (len(columns) - len(row)))
        if not rows or list(rows[0]) != list(columns):
            raise KernelError("header_mismatch")
        series = []
        for index, (name, kind) in enumerate(columns.items()):
            values = [row[index] for row in rows[1:]]
            if kind == "decimal":
                series.append(
                    decimal_series(name, values, precision=precision, scale=scale)
                )
            elif kind == "string" and all(
                value is None or isinstance(value, str) for value in values
            ):
                series.append(pl.Series(name, values, dtype=pl.String))
            elif kind == "boolean" and all(
                value is None or isinstance(value, bool) for value in values
            ):
                series.append(pl.Series(name, values, dtype=pl.Boolean))
            else:
                raise KernelError("unsupported_cell_type")
        return pl.DataFrame(series)
    finally:
        book.close()


def write_table(frame: pl.DataFrame, *, sheet: str) -> bytes:
    """Nouveau classeur simple ; Decimal est du texte exact, relu avant retour.

    Ce résultat en mémoire n'est ni publié, ni un CommitReceipt produit.
    """
    if not frame.width or (frame.height + 1) * frame.width > 100000:
        raise KernelError("cell_limit")
    if not sheet or len(sheet) > 31 or re.search(r"[\\/*?:\[\]]", sheet):
        raise KernelError("sheet_name")
    rows = [tuple(frame.columns)]
    for row in frame.iter_rows():
        converted = []
        for value in row:
            if isinstance(value, Decimal):
                if not value.is_finite():
                    raise KernelError("decimal_invalid")
                value = format(value, "f")
            if value is not None and not isinstance(value, (str, bool, int)):
                raise KernelError("unsupported_cell_type")
            if isinstance(value, str) and len(value) > 32767:
                raise KernelError("text_length")
            converted.append(value)
        rows.append(tuple(converted))
    book = openpyxl.Workbook()
    page = book.active
    write_error = False
    try:
        page.title = sheet
        for r, row in enumerate(rows, 1):
            for c, value in enumerate(row, 1):
                cell = page.cell(r, c, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
        stream = BytesIO()
        book.save(stream)
        result = stream.getvalue()
    except (IllegalCharacterError, ValueError, TypeError, OverflowError):
        write_error = True
    finally:
        book.close()
    # Raise outside the except block: no sensitive library exception is retained.
    if write_error:
        raise KernelError("invalid_output_value")
    check = openpyxl.load_workbook(BytesIO(result), read_only=True, data_only=False)
    try:
        page = check[sheet]
        actual = [tuple(cell.value for cell in row) for row in page.iter_rows()]
        if actual != rows or any(cell.data_type == "f" for row in page for cell in row):
            raise KernelError("readback_mismatch")
    finally:
        check.close()
    return result
